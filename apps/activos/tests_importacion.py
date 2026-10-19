import io

from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.usuarios.models import Usuario

from . import importacion
from .models import EstadoCuraduria, Nivel, NodoActivo, Origen


def excel(filas, encabezado=importacion.COLUMNAS):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Árbol"
    hoja.append(encabezado)
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


class PlantillaTests(TestCase):
    def test_plantilla_tiene_columnas_e_instrucciones(self):
        libro = load_workbook(io.BytesIO(importacion.generar_plantilla()))
        self.assertEqual(libro.sheetnames, ["Árbol", "Instrucciones"])
        self.assertEqual([c.value for c in libro["Árbol"][1]], importacion.COLUMNAS)

    def test_la_plantilla_misma_se_puede_importar(self):
        r = importacion.analizar(io.BytesIO(importacion.generar_plantilla()))
        self.assertEqual(r.errores, [])
        self.assertEqual(r.total_nuevos, 7)


class AnalisisTests(TestCase):
    def setUp(self):
        self.usuario = Usuario.objects.create_user("admin2")

    def test_rutas_validas_y_conteo(self):
        r = importacion.analizar(excel([
            ["Laminadora 6", "Corte", "Hidráulico", "Bomba", "Sello"],
            ["Laminadora 6", "Corte", "Eléctrico"],
            ["LAMINADORA 6", "corte"],  # repetida con otra escritura: no cuenta doble
        ]))
        self.assertEqual(r.errores, [])
        self.assertEqual(r.nuevos, {1: 1, 2: 1, 3: 2, 4: 1, 5: 1})

    def test_errores_por_fila(self):
        r = importacion.analizar(excel([
            ["", "Corte"],
            ["M1", "", "Sistema"],
            ["M1", "x" * 201],
            ["M1", "---"],
            [],
        ]))
        self.assertEqual([fila for fila, _ in r.errores], [2, 3, 4, 5])
        self.assertIn("Falta la máquina", r.errores[0][1])
        self.assertIn("Unidad", r.errores[1][1])

    def test_encabezado_incorrecto(self):
        r = importacion.analizar(excel([["M1"]], encabezado=["Equipo", "Otra"]))
        self.assertEqual(len(r.errores), 1)

    def test_archivo_que_no_es_excel(self):
        r = importacion.analizar(io.BytesIO(b"no soy un excel"))
        self.assertIn("no es un Excel", r.errores[0][1])

    def test_numeros_como_nombres(self):
        r = importacion.analizar(excel([["Laminadora", 6.0]]))
        self.assertEqual(r.rutas[0][1], ["Laminadora", "6"])

    def test_avisa_nombres_parecidos_pero_no_numeros_distintos(self):
        r = importacion.analizar(excel([
            ["M1", "Bomba 1"], ["M1", "Bomba1"], ["M1", "Bomba 2"],
        ]))
        self.assertEqual(len(r.avisos), 1)
        self.assertIn("Bomba1", r.avisos[0])

    def test_avisa_parecido_con_nodo_existente(self):
        NodoActivo.objects.create(nombre="Motor principal", nivel=Nivel.MAQUINA)
        r = importacion.analizar(excel([["Motor principl"]]))
        self.assertEqual(len(r.avisos), 1)


class ImportacionTests(TestCase):
    def setUp(self):
        self.usuario = Usuario.objects.create_user("admin2")
        self.existente = NodoActivo.objects.create(nombre="Laminadora 6", nivel=Nivel.MAQUINA)

    def test_importa_reutilizando_existentes(self):
        archivo = excel([["laminadora 6", "Corte", "Hidráulico"], ["Laminadora 6", "Corte", "Eléctrico"]])
        r = importacion.analizar(archivo)
        self.assertEqual((r.total_nuevos, r.total_existentes), (3, 1))
        self.assertEqual(importacion.importar(r, self.usuario), 3)
        corte = NodoActivo.objects.get(nombre="Corte")
        self.assertEqual(corte.padre, self.existente)
        self.assertEqual((corte.origen, corte.estado_curaduria, corte.creado_por),
                         (Origen.EXCEL, EstadoCuraduria.APROBADO, self.usuario))
        self.assertEqual(corte.hijos.count(), 2)
        self.assertEqual(corte.history.first().history_change_reason, "Importación desde Excel")

    def test_importar_dos_veces_no_duplica(self):
        filas = [["Laminadora 6", "Corte", "Hidráulico"]]
        importacion.importar(importacion.analizar(excel(filas)), self.usuario)
        r = importacion.analizar(excel(filas))
        self.assertEqual(r.total_nuevos, 0)
        self.assertEqual(importacion.importar(r, self.usuario), 0)
        self.assertEqual(NodoActivo.objects.count(), 3)

    def test_con_errores_no_importa_nada(self):
        r = importacion.analizar(excel([["Nueva", "Unidad"], ["", "Huérfana"]]))
        with self.assertRaises(ValueError):
            importacion.importar(r, self.usuario)
        self.assertEqual(NodoActivo.objects.count(), 1)


class AdminImportacionTests(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_superuser("jefe", password="x")
        self.client.force_login(self.admin)

    def test_recorrido_completo(self):
        r = self.client.get("/admin/activos/nodoactivo/plantilla/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("plantilla_arbol_activos.xlsx", r["Content-Disposition"])

        archivo = excel([["Laminadora 6", "Corte", "Hidráulico"]])
        archivo.name = "arbol.xlsx"
        r = self.client.post("/admin/activos/nodoactivo/importar/", {"archivo": archivo})
        self.assertContains(r, "Importar 3 nodos")
        self.assertEqual(NodoActivo.objects.count(), 0)  # Revisar no guarda nada.

        r = self.client.post("/admin/activos/nodoactivo/importar/", {"token": r.context["token"]}, follow=True)
        self.assertContains(r, "3 nodos nuevos")
        self.assertContains(r, "Hidráulico")
        self.assertEqual(NodoActivo.objects.count(), 3)

    def test_archivo_con_errores_muestra_filas(self):
        archivo = excel([["", "Huérfana"]])
        archivo.name = "malo.xlsx"
        r = self.client.post("/admin/activos/nodoactivo/importar/", {"archivo": archivo})
        self.assertContains(r, "No se importó nada")
        self.assertContains(r, "Falta la máquina")
        self.assertNotContains(r, "Importar 0 nodos")

    def test_token_invalido(self):
        r = self.client.post("/admin/activos/nodoactivo/importar/", {"token": "../../x"}, follow=True)
        self.assertContains(r, "ya no está disponible")

    def test_lista_tiene_botones(self):
        r = self.client.get("/admin/activos/nodoactivo/")
        for texto in ("Ver árbol", "Descargar plantilla", "Importar Excel"):
            self.assertContains(r, texto)
