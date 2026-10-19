"""Plantilla de Excel e importación del árbol de activos.

Cada fila del Excel es una ruta: Máquina, Unidad, Sistema, Equipo, Componente.
Primero se analiza el archivo sin guardar nada; solo se importa si no hay errores.
"""

import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz

from .models import EstadoCuraduria, Nivel, NodoActivo, Origen
from .normalizar import normalizar

COLUMNAS = [n.label for n in Nivel]
LARGO_MAXIMO = NodoActivo._meta.get_field("nombre").max_length
SIMILITUD_AVISO = 88  # % de rapidfuzz para avisar de nombres parecidos entre hermanos


# Plantilla

def generar_plantilla():
    """Devuelve los bytes de un .xlsx con la hoja del árbol y una hoja de instrucciones."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Árbol"
    hoja.append(COLUMNAS)
    hoja.append(["Laminadora 6", "Corte", "Hidráulico", "Bomba principal", "Sello mecánico"])
    hoja.append(["Laminadora 6", "Corte", "Eléctrico", "Motor 1"])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E293B")
        celda.alignment = Alignment(horizontal="center")
    for letra in "ABCDE":
        hoja.column_dimensions[letra].width = 28
    hoja.freeze_panes = "A2"

    ayuda = libro.create_sheet("Instrucciones")
    for linea in [
        "Cómo llenar la hoja «Árbol»",
        "",
        "1. Cada fila es una ruta completa: Máquina → Unidad → Sistema → Equipo → Componente.",
        "2. Las columnas del final pueden quedar vacías (por ejemplo, una ruta que llega solo hasta Equipo).",
        "3. No deje huecos en medio: si hay Sistema, debe haber Unidad y Máquina.",
        "4. Repita el nombre de la máquina, unidad, etc. en cada fila; no combine celdas.",
        "5. Mayúsculas, tildes y signos no importan para reconocer nodos existentes.",
        f"6. Cada nombre tiene máximo {LARGO_MAXIMO} caracteres.",
        "7. Borre las filas de ejemplo antes de importar.",
    ]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=14)
    ayuda.column_dimensions["A"].width = 100

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# Análisis

@dataclass
class Resultado:
    rutas: list = field(default_factory=list)       # (n.º de fila, [nombres]) válidas
    errores: list = field(default_factory=list)     # (n.º de fila, motivo)
    avisos: list = field(default_factory=list)      # textos
    nuevos: dict = field(default_factory=lambda: {n.value: 0 for n in Nivel})
    existentes: dict = field(default_factory=lambda: {n.value: 0 for n in Nivel})

    @property
    def total_nuevos(self):
        return sum(self.nuevos.values())

    @property
    def total_existentes(self):
        return sum(self.existentes.values())

    @property
    def por_nivel(self):
        return [(n.label, self.nuevos[n.value], self.existentes[n.value]) for n in Nivel]


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", " ", str(valor)).strip()


def _leer_filas(archivo):
    """Devuelve [(n.º de fila, [5 textos])], o lanza ValueError si el archivo no sirve."""
    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ValueError("El archivo no es un Excel válido (.xlsx).")
    hoja = libro["Árbol"] if "Árbol" in libro.sheetnames else libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezado = [normalizar(_texto(v)) for v in (next(filas, None) or [])][: len(COLUMNAS)]
    if encabezado != [normalizar(c) for c in COLUMNAS]:
        raise ValueError("La primera fila debe tener las columnas: " + ", ".join(COLUMNAS) + ".")
    resultado = []
    for numero, valores in enumerate(filas, start=2):
        textos = [_texto(v) for v in list(valores or [])[: len(COLUMNAS)]]
        textos += [""] * (len(COLUMNAS) - len(textos))
        resultado.append((numero, textos))
    libro.close()
    return resultado


def _parecidos(a, b):
    """Nombres normalizados distintos pero casi iguales, con los mismos números ('bomba 1' / 'bomba1')."""
    if a == b or re.findall(r"\d+", a) != re.findall(r"\d+", b):
        return False
    return a.replace(" ", "") == b.replace(" ", "") or fuzz.ratio(a, b) >= SIMILITUD_AVISO


class _Arbol:
    """Recorre las rutas sobre el árbol actual, distinguiendo nodos existentes y nuevos."""

    def __init__(self):
        self.existentes = {}  # (padre_id, nombre_normalizado) → NodoActivo
        self.hijos_reales = {}  # padre_id → [nombre_normalizado]
        for nodo in NodoActivo.objects.filter(is_active=True):
            self.existentes[(nodo.padre_id, nodo.nombre_normalizado)] = nodo
            self.hijos_reales.setdefault(nodo.padre_id, []).append(nodo.nombre_normalizado)
        self.nuevos = {}  # clave de la ruta normalizada → nombre original

    def recorrer(self, nombres):
        """Genera (nivel, clave_padre, padre_real, nombre, normalizado, existente) por cada nivel de la ruta."""
        padre_real, clave = None, ()
        for nivel, nombre in enumerate(nombres, start=1):
            normalizado = normalizar(nombre)
            clave_padre, clave = clave, clave + (normalizado,)
            existente = None
            if padre_real is not None or nivel == Nivel.MAQUINA:
                existente = self.existentes.get((padre_real.pk if padre_real else None, normalizado))
            yield nivel, clave_padre, padre_real, nombre, normalizado, existente
            padre_real = existente


def analizar(archivo):
    resultado = Resultado()
    try:
        filas = _leer_filas(archivo)
    except ValueError as error:
        resultado.errores.append((1, str(error)))
        return resultado

    for numero, textos in filas:
        if not any(textos):
            continue
        ultimo = max(i for i, t in enumerate(textos) if t)
        huecos = [COLUMNAS[i] for i in range(ultimo) if not textos[i]]
        largos = [COLUMNAS[i] for i in range(ultimo + 1) if len(textos[i]) > LARGO_MAXIMO]
        if not textos[0]:
            resultado.errores.append((numero, "Falta la máquina."))
        elif huecos:
            resultado.errores.append((numero, "Hay columnas vacías en medio de la ruta: " + ", ".join(huecos) + "."))
        elif largos:
            resultado.errores.append((numero, f"Nombre de más de {LARGO_MAXIMO} caracteres en: " + ", ".join(largos) + "."))
        elif any(not normalizar(t) for t in textos[: ultimo + 1]):
            resultado.errores.append((numero, "Hay un nombre sin letras ni números."))
        else:
            resultado.rutas.append((numero, textos[: ultimo + 1]))

    if not resultado.rutas and not resultado.errores:
        resultado.errores.append((1, "El archivo no tiene filas con datos."))

    arbol = _Arbol()
    contados = set()
    hermanos_nuevos = {}  # clave del padre → [(normalizado, nombre)]
    for _, nombres in resultado.rutas:
        for nivel, clave_padre, padre_real, nombre, normalizado, existente in arbol.recorrer(nombres):
            clave = clave_padre + (normalizado,)
            if clave in contados:
                continue
            contados.add(clave)
            if existente:
                resultado.existentes[nivel] += 1
                continue
            resultado.nuevos[nivel] += 1
            hermanos = hermanos_nuevos.setdefault(clave_padre, [])
            otros = [(n, n) for n in arbol.hijos_reales.get(padre_real.pk if padre_real else None, [])
                     if padre_real or nivel == Nivel.MAQUINA] + hermanos
            for otro_normalizado, otro_nombre in otros:
                if _parecidos(normalizado, otro_normalizado):
                    lugar = " › ".join(nombres[: nivel - 1]) or "las máquinas"
                    resultado.avisos.append(f"«{nombre}» se parece a «{otro_nombre}» en {lugar}.")
            hermanos.append((normalizado, nombre))
    return resultado


# Importación

@transaction.atomic
def importar(resultado, usuario):
    """Crea los nodos nuevos. Solo se llama si el análisis no tuvo errores."""
    if resultado.errores:
        raise ValueError("No se importa un archivo con errores.")
    arbol = _Arbol()
    creados = {}  # clave de la ruta normalizada → NodoActivo creado en esta importación
    for _, nombres in resultado.rutas:
        padre = None
        clave = ()
        for nivel, nombre in enumerate(nombres, start=1):
            normalizado = normalizar(nombre)
            clave += (normalizado,)
            nodo = creados.get(clave) or arbol.existentes.get((padre.pk if padre else None, normalizado))
            if nodo is None:
                nodo = NodoActivo(nombre=nombre, nivel=nivel, padre=padre, origen=Origen.EXCEL,
                                  estado_curaduria=EstadoCuraduria.APROBADO, creado_por=usuario)
                nodo._change_reason = "Importación desde Excel"
                nodo.save()
                creados[clave] = nodo
            padre = nodo
    return len(creados)
